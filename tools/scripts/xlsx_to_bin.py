#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
策划表xlsx转加密bin工具
支持将xlsx文件转换为加密的二进制格式, 供C++程序读取
"""

import os
import sys
import struct
import hashlib
import openpyxl
from openpyxl import load_workbook
import json
import argparse
from typing import Dict, List, Any, Tuple
import zlib
from encrypt import create_encryptor, EncryptionType

class XlsxToBinConverter:
    """Xlsx到Bin的转换器"""

    # 文件头标识 "CHGD" (Chen Game Data)
    FILE_HEADER = b"CHGD"
    # 版本号
    VERSION = 2  # 更新版本号以支持新的加密方案

    # 简单的XOR加密密钥（保留用于兼容）
    ENCRYPT_KEY = 0xABCD1234

    def __init__(self, encryption_type: str = "aes256_gcm", password: str = None):
        self.tables = {}  # 存储所有表数据
        self.encryption_type = encryption_type
        self.password = password or "chen_default_password"
        self.encryptor = create_encryptor(encryption_type)

    def _encrypt_data(self, data: bytes) -> bytes:
        """使用配置的加密方案加密数据"""
        if self.encryption_type == "xor":
            # 保留兼容性
            return self.encryptor.encrypt(data, key=self.ENCRYPT_KEY.to_bytes(4, 'big'))['data']
        else:
            # 使用新的安全加密方案
            encrypted = self.encryptor.encrypt(data, self.password)
            # 将加密类型和元数据序列化到数据前
            metadata = encrypted['metadata']
            metadata_bytes = json.dumps({
                'type': int(encrypted['type']),
                'metadata': {k: v.hex() if isinstance(v, bytes) else v for k, v in metadata.items()}
            }).encode('utf-8')
            return struct.pack('<I', len(metadata_bytes)) + metadata_bytes + encrypted['data']

    def _compress_data(self, data: bytes) -> bytes:
        """压缩数据"""
        return zlib.compress(data, level=9)

    def _read_xlsx(self, file_path: str) -> Dict[str, Any]:
        """读取xlsx文件"""
        try:
            wb = load_workbook(file_path, read_only=True)

            # 获取第一个工作表
            ws = wb.active

            # 读取表头（第一行）
            headers = []
            for cell in ws[1]:
                if cell.value:
                    headers.append(str(cell.value))
                else:
                    break

            if not headers:
                raise ValueError("No headers found in xlsx file")

            # 读取数据行
            data_rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_data = {}
                for i, value in enumerate(row):
                    if i < len(headers):
                        # 处理不同类型的数据
                        if value is None:
                            row_data[headers[i]] = None
                        elif isinstance(value, (int, float)):
                            row_data[headers[i]] = value
                        elif isinstance(value, str):
                            # 尝试转换为数字
                            try:
                                if '.' in value:
                                    row_data[headers[i]] = float(value)
                                else:
                                    row_data[headers[i]] = int(value)
                            except ValueError:
                                row_data[headers[i]] = value
                        else:
                            row_data[headers[i]] = str(value)

                # 只添加非空行
                if any(v is not None for v in row_data.values()):
                    data_rows.append(row_data)

            wb.close()

            return {
                'headers': headers,
                'data': data_rows,
                'row_count': len(data_rows)
            }

        except Exception as e:
            raise Exception(f"Failed to read xlsx file {file_path}: {str(e)}")

    def _serialize_table(self, table_name: str, table_data: Dict[str, Any]) -> bytes:
        """序列化表数据为二进制格式"""
        headers = table_data['headers']
        data = table_data['data']

        # 构建表数据
        table_bytes = bytearray()

        # 表名长度和表名
        table_name_bytes = table_name.encode('utf-8')
        table_bytes.extend(struct.pack('<I', len(table_name_bytes)))
        table_bytes.extend(table_name_bytes)

        # 字段数量
        table_bytes.extend(struct.pack('<I', len(headers)))

        # 每个字段名
        for header in headers:
            header_bytes = header.encode('utf-8')
            table_bytes.extend(struct.pack('<I', len(header_bytes)))
            table_bytes.extend(header_bytes)

        # 行数
        table_bytes.extend(struct.pack('<I', len(data)))

        # 序列化每一行
        for row in data:
            # 为每个字段写入数据
            for header in headers:
                value = row.get(header)

                if value is None:
                    # NULL值用特殊标记
                    table_bytes.extend(struct.pack('<B', 0))
                elif isinstance(value, int):
                    # 整数类型 - 统一使用int32
                    table_bytes.extend(struct.pack('<Bi', 3, value))
                elif isinstance(value, float):
                    # 浮点类型
                    table_bytes.extend(struct.pack('<Bd', 5, value))
                elif isinstance(value, str):
                    # 字符串类型
                    str_bytes = value.encode('utf-8')
                    str_len = len(str_bytes)
                    if str_len < 255:
                        # 短字符串
                        table_bytes.extend(struct.pack('<BB', 6, str_len))
                        table_bytes.extend(str_bytes)
                    else:
                        # 长字符串
                        table_bytes.extend(struct.pack('<BI', 7, str_len))
                        table_bytes.extend(str_bytes)
                else:
                    # 其他类型转为字符串
                    str_value = str(value)
                    str_bytes = str_value.encode('utf-8')
                    if len(str_bytes) < 255:
                        table_bytes.extend(struct.pack('<BB', 6, len(str_bytes)))
                    else:
                        table_bytes.extend(struct.pack('<BI', 7, len(str_bytes)))
                    table_bytes.extend(str_bytes)

        return bytes(table_bytes)

    def convert_to_bin(self, xlsx_files: List[str], output_file: str):
        """转换xlsx文件为加密的bin文件"""
        print(f"开始转换 {len(xlsx_files)} 个xlsx文件...")
        print(f"使用加密方案: {self.encryption_type}")

        # 如果是多个文件且输出是目录，或输出路径包含通配符，使用分离模式
        is_output_dir = output_file.endswith('/') or output_file.endswith('\\') or os.path.isdir(output_file)

        if (len(xlsx_files) > 1 and is_output_dir) or '*' in output_file:
            # 分离模式：每个xlsx转换为单独的bin文件
            self._convert_separate(xlsx_files, output_file)
        else:
            # 合并模式：所有xlsx合并到一个bin文件
            self._convert_merged(xlsx_files, output_file)

    def _convert_separate(self, xlsx_files: List[str], output_pattern: str):
        """分离模式：每个xlsx转换为单独的bin文件"""
        print("使用分离模式：每个文件单独转换")

        # 创建输出目录
        output_dir = os.path.dirname(output_pattern)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        for xlsx_file in xlsx_files:
            if not os.path.exists(xlsx_file):
                print(f"警告: 文件不存在 {xlsx_file}")
                continue

            table_name = os.path.splitext(os.path.basename(xlsx_file))[0]
            print(f"\n处理文件: {xlsx_file}")

            # 读取单个文件
            table_data = self._read_xlsx(xlsx_file)

            # 生成输出文件名
            if '*' in output_pattern:
                # 替换通配符
                output_file = output_pattern.replace('*', table_name)
            else:
                # 使用输出目录 + 文件名
                output_file = os.path.join(output_dir, f"{table_name}.bin")

            # 转换单个表
            self._convert_single_table(table_name, table_data, output_file)

    def _convert_merged(self, xlsx_files: List[str], output_file: str):
        """合并模式：所有xlsx合并到一个bin文件"""
        print("使用合并模式：所有文件合并到一个bin")

        # 创建输出目录
        output_dir = os.path.dirname(output_file)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        # 读取所有xlsx文件
        for xlsx_file in xlsx_files:
            if not os.path.exists(xlsx_file):
                print(f"警告: 文件不存在 {xlsx_file}")
                continue

            table_name = os.path.splitext(os.path.basename(xlsx_file))[0]
            print(f"读取文件: {xlsx_file} -> 表名: {table_name}")

            table_data = self._read_xlsx(xlsx_file)
            self.tables[table_name] = table_data

            print(f"  - 字段数: {len(table_data['headers'])}")
            print(f"  - 行数: {table_data['row_count']}")

        # 序列化所有表
        all_tables_bytes = bytearray()

        for table_name, table_data in self.tables.items():
            table_bytes = self._serialize_table(table_name, table_data)
            # 表长度 + 表数据
            all_tables_bytes.extend(struct.pack('<I', len(table_bytes)))
            all_tables_bytes.extend(table_bytes)

        # 构建文件头
        file_header = bytearray()
        # 文件标识
        file_header.extend(self.FILE_HEADER)
        # 版本号
        file_header.extend(struct.pack('<I', self.VERSION))
        # 表数量
        file_header.extend(struct.pack('<I', len(self.tables)))
        # 保留字段
        file_header.extend(struct.pack('<I', 0))

        # 合并所有数据
        all_data = bytes(file_header) + bytes(all_tables_bytes)

        # 压缩数据
        compressed_data = self._compress_data(all_data)

        # 计算校验和
        checksum = hashlib.md5(compressed_data).digest()

        # 最终文件: 校验和 + 压缩数据
        final_data = checksum + compressed_data

        # 加密
        encrypted_data = self._encrypt_data(final_data)

        # 写入文件
        with open(output_file, 'wb') as f:
            f.write(encrypted_data)

        print(f"\n转换完成! 输出文件: {output_file}")
        print(f"包含表数量: {len(self.tables)}")
        print(f"原始大小: {len(all_data)} 字节")
        print(f"压缩后: {len(compressed_data)} 字节")
        print(f"压缩率: {len(compressed_data) / len(all_data) * 100:.1f}%")

    def _convert_single_table(self, table_name: str, table_data: Dict[str, Any], output_file: str):
        """转换单个表为bin文件"""
        print(f"  - 字段数: {len(table_data['headers'])}")
        print(f"  - 行数: {table_data['row_count']}")

        # 序列化表数据
        table_bytes = self._serialize_table(table_name, table_data)

        # 构建文件头
        file_header = bytearray()
        # 文件标识
        file_header.extend(self.FILE_HEADER)
        # 版本号
        file_header.extend(struct.pack('<I', self.VERSION))
        # 表数量
        file_header.extend(struct.pack('<I', 1))
        # 保留字段
        file_header.extend(struct.pack('<I', 0))

        # 表长度 + 表数据
        all_tables_bytes = bytearray()
        all_tables_bytes.extend(struct.pack('<I', len(table_bytes)))
        all_tables_bytes.extend(table_bytes)

        # 合并所有数据
        all_data = bytes(file_header) + bytes(all_tables_bytes)

        # 压缩数据
        compressed_data = self._compress_data(all_data)

        # 计算校验和
        checksum = hashlib.md5(compressed_data).digest()

        # 最终文件: 校验和 + 压缩数据
        final_data = checksum + compressed_data

        # 加密
        encrypted_data = self._encrypt_data(final_data)

        # 写入文件
        with open(output_file, 'wb') as f:
            f.write(encrypted_data)

        print(f"  -> 输出文件: {output_file}")
        print(f"  原始大小: {len(all_data)} 字节, 压缩后: {len(compressed_data)} 字节")

def main():
    parser = argparse.ArgumentParser(description='Xlsx到加密Bin文件转换工具')
    parser.add_argument('input', nargs='+', help='输入的xlsx文件路径')
    parser.add_argument('-o', '--output', required=True, help='输出的bin文件路径')
    parser.add_argument('-v', '--verbose', action='store_true', help='显示详细信息')
    parser.add_argument('-e', '--encrypt', default='aes256_gcm',
                       choices=['none', 'xor', 'aes256_gcm', 'chacha20_poly1305', 'custom'],
                       help='加密方案 (默认: aes256_gcm)')
    parser.add_argument('-p', '--password', help='加密密码 (默认使用默认密码)')

    args = parser.parse_args()

    converter = XlsxToBinConverter(encryption_type=args.encrypt, password=args.password)

    try:
        converter.convert_to_bin(args.input, args.output)
    except Exception as e:
        print(f"错误: {str(e)}")
        sys.exit(1)

if __name__ == '__main__':
    main()