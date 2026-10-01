#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CSV转XLSX工具
支持将CSV文件转换为XLSX格式，方便后续使用xlsx_to_bin工具处理
"""

import os
import sys
import csv
import argparse
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import chardet

class CsvToXlsxConverter:
    """CSV到XLSX的转换器"""

    def __init__(self, auto_detect_encoding: bool = True, default_encoding: str = 'utf-8'):
        self.auto_detect_encoding = auto_detect_encoding
        self.default_encoding = default_encoding

    def _detect_encoding(self, file_path: str) -> str:
        """自动检测文件编码"""
        if not self.auto_detect_encoding:
            return self.default_encoding

        with open(file_path, 'rb') as f:
            raw_data = f.read(10000)  # 读取前10KB用于检测
            result = chardet.detect(raw_data)
            encoding = result['encoding']
            confidence = result['confidence']

            print(f"检测到编码: {encoding} (置信度: {confidence:.2f})")

            # 如果置信度太低，使用默认编码
            if confidence < 0.7:
                print(f"置信度较低，使用默认编码: {self.default_encoding}")
                return self.default_encoding

            return encoding or self.default_encoding

    def _detect_delimiter(self, file_path: str, encoding: str) -> str:
        """自动检测CSV分隔符"""
        with open(file_path, 'r', encoding=encoding, newline='') as f:
            # 读取第一行
            first_line = f.readline()

            # 尝试不同的分隔符
            delimiters = [',', ';', '\t', '|']
            delimiter_counts = {}

            for delimiter in delimiters:
                count = first_line.count(delimiter)
                if count > 0:
                    delimiter_counts[delimiter] = count

            if delimiter_counts:
                # 选择出现次数最多的分隔符
                detected_delimiter = max(delimiter_counts, key=delimiter_counts.get)
                print(f"检测到分隔符: '{detected_delimiter}'")
                return detected_delimiter

            return ','  # 默认使用逗号

    def _read_csv_data(self, file_path: str) -> Dict[str, Any]:
        """读取CSV文件数据"""
        encoding = self._detect_encoding(file_path)
        delimiter = self._detect_delimiter(file_path, encoding)

        try:
            with open(file_path, 'r', encoding=encoding, newline='') as f:
                # 使用csv模块读取数据
                csv_reader = csv.reader(f, delimiter=delimiter)

                # 读取所有行
                rows = list(csv_reader)

                if not rows:
                    raise ValueError("CSV文件为空")

                # 第一行作为表头
                headers = rows[0]
                data_rows = rows[1:]

                # 清理表头（去除空字符串）
                headers = [h.strip() for h in headers if h.strip()]

                # 清理数据行
                cleaned_data = []
                for row in data_rows:
                    # 确保行的长度与表头一致
                    while len(row) < len(headers):
                        row.append('')
                    row_data = {}
                    for i, value in enumerate(row[:len(headers)]):
                        if i < len(headers):
                            # 尝试转换数据类型
                            cleaned_value = self._convert_value(value.strip())
                            row_data[headers[i]] = cleaned_value

                    # 只添加非空行
                    if any(v is not None and v != '' for v in row_data.values()):
                        cleaned_data.append(row_data)

                return {
                    'headers': headers,
                    'data': cleaned_data,
                    'row_count': len(cleaned_data)
                }

        except Exception as e:
            raise Exception(f"读取CSV文件失败 {file_path}: {str(e)}")

    def _convert_value(self, value: str) -> Any:
        """尝试转换字符串值为适当的数据类型"""
        if not value or value.lower() in ('null', 'none', 'nan', ''):
            return None

        # 尝试转换为整数
        try:
            if '.' not in value and 'e' not in value.lower():
                return int(value)
        except ValueError:
            pass

        # 尝试转换为浮点数
        try:
            return float(value)
        except ValueError:
            pass

        # 尝试转换为布尔值
        if value.lower() in ('true', 'false'):
            return value.lower() == 'true'

        # 返回字符串
        return value

    def _create_xlsx(self, csv_data: Dict[str, Any], output_path: str,
                    auto_format: bool = True) -> None:
        """创建XLSX文件"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

        headers = csv_data['headers']
        data = csv_data['data']

        # 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            if auto_format:
                # 设置表头样式
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # 写入数据行
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, '')
                ws.cell(row=row_idx, column=col_idx, value=value)

        # 自动调整列宽
        if auto_format:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter

                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass

                # 设置列宽（最小10，最大50）
                adjusted_width = min(max(max_length + 2, 10), 50)
                ws.column_dimensions[column_letter].width = adjusted_width

        # 保存文件
        wb.save(output_path)

    def convert_csv_to_xlsx(self, csv_files: List[str], output_pattern: str,
                           auto_format: bool = True) -> None:
        """转换CSV文件为XLSX格式"""
        print(f"开始转换 {len(csv_files)} 个CSV文件...")

        # 检查输出是否为目录
        is_output_dir = output_pattern.endswith('/') or output_pattern.endswith('\\') or os.path.isdir(output_pattern)

        # 创建输出目录
        if is_output_dir:
            output_dir = output_pattern if os.path.isdir(output_pattern) else output_pattern
            if not os.path.exists(output_dir):
                os.makedirs(output_dir)
        else:
            output_dir = os.path.dirname(output_pattern)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir)

        converted_files = []

        for csv_file in csv_files:
            if not os.path.exists(csv_file):
                print(f"警告: 文件不存在 {csv_file}")
                continue

            print(f"\n处理文件: {csv_file}")

            try:
                # 读取CSV数据
                csv_data = self._read_csv_data(csv_file)
                print(f"  - 字段数: {len(csv_data['headers'])}")
                print(f"  - 数据行数: {csv_data['row_count']}")

                # 生成输出文件名
                if '*' in output_pattern:
                    # 替换通配符
                    base_name = os.path.splitext(os.path.basename(csv_file))[0]
                    output_file = output_pattern.replace('*', base_name)
                elif is_output_dir:
                    # 使用输出目录 + 文件名
                    base_name = os.path.splitext(os.path.basename(csv_file))[0]
                    output_file = os.path.join(output_dir, f"{base_name}.xlsx")
                else:
                    # 直接使用指定的输出文件（仅当输入只有一个文件时）
                    if len(csv_files) == 1:
                        output_file = output_pattern
                    else:
                        base_name = os.path.splitext(os.path.basename(csv_file))[0]
                        output_file = os.path.join(output_dir, f"{base_name}.xlsx")

                # 转换为XLSX
                self._create_xlsx(csv_data, output_file, auto_format)
                print(f"  -> 输出文件: {output_file}")

                converted_files.append(output_file)

            except Exception as e:
                print(f"  错误: {str(e)}")
                continue

        print(f"\n转换完成! 共转换 {len(converted_files)} 个文件")

        if converted_files:
            print("输出文件:")
            for file in converted_files:
                print(f"  - {file}")

def main():
    parser = argparse.ArgumentParser(description='CSV到XLSX转换工具')
    parser.add_argument('input', nargs='+', help='输入的CSV文件路径')
    parser.add_argument('-o', '--output', required=True,
                       help='输出的XLSX文件路径（可使用*通配符或指定目录）')
    parser.add_argument('-v', '--verbose', action='store_true', help='显示详细信息')
    parser.add_argument('--no-auto-format', action='store_true',
                       help='禁用自动格式化（表头样式、列宽调整等）')
    parser.add_argument('--encoding', default='utf-8',
                       help='指定文件编码（默认: utf-8）')
    parser.add_argument('--no-auto-encoding', action='store_true',
                       help='禁用自动编码检测')

    args = parser.parse_args()

    converter = CsvToXlsxConverter(
        auto_detect_encoding=not args.no_auto_encoding,
        default_encoding=args.encoding
    )

    try:
        converter.convert_csv_to_xlsx(
            args.input,
            args.output,
            auto_format=not args.no_auto_format
        )
    except Exception as e:
        print(f"错误: {str(e)}")
        sys.exit(1)

if __name__ == '__main__':
    main()